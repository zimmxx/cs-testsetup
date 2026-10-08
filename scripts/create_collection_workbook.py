import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
catalog=json.loads((ROOT/'.local/workbook-catalog.json').read_text(encoding='utf-8-sig'))
doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.65)
section.left_margin=section.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(13)
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(7)
header=section.header.paragraphs[0];header.text='CORNERSTONE   |   Testing setup collection workbook';header.runs[0].font.size=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Version 1.0   |   30 September 2026   |   Page ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text='',style=None):return doc.add_paragraph(text,style)
def heading(text):doc.add_heading(text,1)
def page(text):doc.add_page_break();heading(text)
def h2(text):doc.add_heading(text,2)
def field(label,hint='Enter value or mark unknown',lines=1):
    para=p();para.add_run(label+'  ').bold=True
    run=para.add_run('['+hint+']');run.italic=True;run.font.color.rgb=RGBColor.from_string('777777')
    for _ in range(lines-1):p('________________________________________________________________________')
def checklist(items):
    for item in items:p('☐  '+item)
def table(headers,rows,widths=None,size=10):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side,val in [('top',95),('bottom',95),('left',105),('right',105)]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
            props.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243650' if ri==0 else ('F3F6FA' if ri%2==0 else 'FFFFFF'));props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.size=Pt(size)
                    if ri==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    p();return t
blank='[Enter / verify]'

doc.add_paragraph('CORNERSTONE testing setup collection workbook','Title')
p('Equipment models measurement setups and work guidelines','Subtitle')
p('Use this editable workbook to collect the evidence needed to update cs-testsetup. Fill in known values, attach or name the source files, and mark unconfirmed details explicitly. Return the completed workbook and its referenced files so the app catalog, directory links, models and guidance can be updated consistently.')
field('Prepared by');field('Lab or team');field('Review owner');field('Workbook revision and date')
h2('How to use this workbook')
p('Complete the master registers first. Copy the equipment sheets for each physical instrument or assembly, and copy the setup and guideline sheets for each real measurement arrangement. Replace bracketed prompts with your findings. Add rows or insert photos where needed; keep equipment IDs and setup IDs stable.')
p('Use these review states throughout: Missing, Collected, Needs verification, Verified, Approved, or Not applicable. Give every important value a source and review date. A model number or service listing alone does not establish the installed configuration or current readiness.')
table(['Priority','Collect first','Review state'],[
('1','Installed equipment identities and actual setup photos',blank),('2','Port to port wiring, fibre models and connector ends',blank),('3','Measured dimensions, vendor CAD and model permissions',blank),('4','Approved work guidelines and test parameter limits',blank),('5','Real timing records and example measurement data',blank)], [0.65,5.0,1.45])
p('This workbook is a collection and review tool. Draft prompts are not approved operating procedures. Keep sensitive or unpublished material outside public folders until the owner approves publication.')

page('Equipment collection register')
p('Starting records from the current app are listed below. Confirm the exact installed part number and add separate records for instruments with different configurations. Use the blank columns to track your evidence.')
table(['Equipment ID','Current record','Exact unit and evidence','State'],[(e['id'],e['name']+'\n'+e['model'],blank,'[State]') for e in catalog['equipment']],[1.02,2.75,2.5,.83],9.5)
field('Additional equipment to add','Include sources, mainframes, adapters, probes, controllers, computers and mechanical assemblies',2)
p('Record serial numbers and asset tags in an access-appropriate inventory. The present app has no user access controls for published library metadata.')

page('Equipment identity and specifications sheet')
p('Duplicate this sheet and the following model sheet for each equipment record. Separate manufacturer ratings from the locally validated operating range.')
field('Equipment ID','Existing stable ID or new lowercase hyphenated ID')
field('Equipment name and purpose');field('Manufacturer model and exact variant');field('Serial number asset tag and quantity','Include publication restriction if applicable')
field('Lab location storage location and owner');field('Condition availability and calibration due date')
table(['Parameter and units','Manufacturer specification','Local validated limit','Evidence and date'],[(v,blank,blank,blank) for v in ['Wavelength / frequency','Power / voltage / current','Sensitivity / noise / resolution','Sweep speed / sample rate','Travel / axes / accuracy','Other critical parameter']],[1.65,1.75,1.8,1.9],10)
table(['Port name','Signal and direction','Connector or interface','Limits and restrictions'],[(blank,blank,blank,blank) for _ in range(3)],[1.4,1.75,1.9,2.05],10)
field('Compatible alternatives and required adapters');field('Firmware driver mainframe and software dependencies')
field('Manufacturer page manual and calibration evidence','Filename or URL and page number')

page('Equipment pictures and detailed model sheet')
field('Equipment ID');field('Physical dimensions and units','Width depth height or dimension drawing; confirm axes and origin')
field('Moving parts and useful interactive features','Travel ranges, rotation, cable bend space, removable modules and selectable parts')
checklist(['Front side rear and top photos show the exact installed variant','Closeups identify all relevant ports and connector faces','At least one known dimension or scale reference is recorded','CAD or dimension drawing matches the variant and revision','Source owner and permission for repository publication are recorded'])
table(['Asset','Source filename or URL','Target relative path','State'],[
('Equipment picture',blank,'library/images/<equipment-id>-front.jpg','[State]'),('Detailed viewer model',blank,'library/models/<equipment-id>.glb','[State]'),('Original CAD and drawings',blank,'library/references/<equipment-id>/','[State]'),('Additional reference photos',blank,'library/references/<equipment-id>/','[State]')],[1.5,2.3,2.45,.85],10)
field('CAD format revision and conversion tool','STEP STP IGES native CAD OBJ etc; final viewer file is GLB 2.0')
field('Model source licence and publication permission',lines=2)
field('Model quality review','Geometry, scale reference, orientation, missing parts, textures and file size',2)
p('Current viewer requirements: self-contained GLB 2.0, embedded textures, up to 40 MB, without external compression decoders. The app scene normalises sizes for illustration. Port snapping, animated mechanisms and clearance calculations are not yet implemented.')
field('Reviewed by date and remaining issues')

page('Measurement setup collection register')
p('These are catalog references and proposed configurations, not a verified booking inventory. Record the actual readiness separately; do not infer installed instrument models from another platform.')
table(['Setup ID and name','Scale and band','Catalog status','Actual evidence and state'],[(s['id']+'\n'+s['name'],s['scale']+' / '+s['mode']+'\n'+(', '.join(s['bands']) or 'No optical band'),s['status'],blank) for s in catalog['setups']],[2.55,1.35,1.35,1.85],9.5)
field('Additional setups and variants to add',lines=2)

page('Actual measurement setup sheet')
p('Duplicate this and the next two sheets for each setup. Use photographs and the physical bench to record the real configuration, including control and observation equipment.')
field('Setup ID name revision and owner');field('Lab location and operating readiness');field('Scale mode and wavelength bands','Chip or wafer; optical electrical or mixed; validated wavelength range')
field('Measurement purpose and supported outputs','e.g. transmission, loss, resistance or phase shift')
field('DUT platform holder and geometry','Chip / wafer size, device map, grating or edge coupling, pad map and temperature control')
field('Setup overview photo or drawing','Insert photo here or give filename; identify camera view and date',2)
table(['Instance label','Equipment ID and actual unit','Role and position','Photo marker'],[(blank,blank,blank,blank) for _ in range(5)],[1.4,2.4,2.25,1.05],10)
field('Interactive marker locations','For a photo record each marker as x and y percent from the top left',2)
field('Measured capabilities limitations and exclusions',lines=2)
field('Required adapters software and missing equipment')

page('Signal connections and fibre sheet')
p('Give every connection its own ID. Connector A belongs to the source cable end and connector B to the destination cable end. Record mating equipment ports and adapters separately. Mixed fibres and connector types are valid records; compatibility still needs checking.')
table(['Link ID and type','From instance and port','To instance and port','Cable or fibre model'],[(blank,blank,blank,blank) for _ in range(5)],[1.25,1.9,1.9,2.05],10)
table(['Link ID','Fibre type and range','Connector A to B','Length adapters and notes'],[('[ID]',blank,blank,blank) for _ in range(5)],[.9,1.9,1.85,2.45],10)
field('Fibre details to verify','Manufacturer part number, SM/PM, wavelength range, polarisation axis, numerical aperture, fibre tip and bend radius',2)
field('Electrical link details','Cable / probe model, contact type, shielding, grounding, compliance and polarity',2)
field('Triggers and control links','USB LAN GPIB TTL synchronisation; distinguish these from measured optical/electrical paths')
checklist(['Each physical connection is traced on the actual bench','Cable end labels match photographs and port specifications','Fibre range covers the full planned sweep','PM orientation and adapter compatibility are checked where relevant','Connection IDs match the exported Build setup JSON'])

page('Measurement parameters and validation sheet')
field('Setup ID measurement recipe version and approver')
table(['Parameter','Value and units','Permitted range or criterion','Evidence'],[(name,blank,blank,blank) for name in ['Wavelength start stop step','Continuous sweep speed / dwell','Input optical power and polarisation','Detector range averaging sample rate','Voltage current and compliance','Temperature setpoint stability','Alignment criteria and reference','Chip / wafer die / device counts','Repeats conditions and scan direction','Data filenames and acquisition format']],[2.25,1.45,2.15,1.25],10)
field('Reference measurement and calibration method',lines=2)
field('Acceptance criteria and quality checks','Saturation, noise, drift, repeatability, fit quality and permitted exclusions',2)
field('Validation result and exceptions','Attach representative traces and the reviewer conclusion',2)
field('Validated by date and next review')

page('Work guideline collection sheet')
field('Guideline title document ID version and owner');field('Applicable setup IDs and user prerequisites')
field('Approved SOP and training references','Exact filenames URLs pages and revision numbers')
p('For each stage, provide the actual approved actions, settings and expected result. Add screenshots or link to the relevant SOP. Include who may perform the task, stop conditions and how to recover from common errors.')
table(['Stage','Approved actions and settings','Expected result and evidence'],[(name,'[Enter steps or SOP section]','[Enter checks / image / file]') for name in ['Preparation and authorisation','Instrument and calibration checks','Load DUT and identify devices','Connect and align the setup','Acquire reference / background','Configure and run measurement','Review and repeat when needed','Shut down and unload','Post process and hand over']],[1.65,3.15,2.3],10)
field('Common errors and approved recovery steps',lines=2)
field('Reviewer approval date and training signoff')
p('The app currently contains draft workflow outlines. Only replace them with operational instructions after the responsible lab reviewer has approved the content.')

page('Timing data and measurement evidence sheet')
field('Setup ID recipe conditions and timing observer')
p('Record observed durations and the basis of each value. Repeat representative jobs to separate fixed preparation time from per-chip, per-die, per-device and per-scan time. Attach raw timing logs when available.')
table(['Activity','Time and units','Basis and sample count','Source / uncertainty'],[(v,blank,blank,blank) for v in ['Session preparation','Load / unload chip or wafer','Move between dies or devices','Alignment / realignment','Reference acquisition','Sweep / electrical sequence','Save transfer or instrument overhead','Post processing','Contingency / interruptions']],[2.25,1.25,1.85,1.75],10)
field('Representative job size and measured total','Chips wafers dies devices repeats conditions; observed start and finish')
field('Planner comparison and adjustment','Estimated time, actual time, explanation of difference',2)
field('Example raw data and processed output','Files, units, column meanings, references and expected analysis results',2)
field('Data retention access and publication restrictions')

page('Directory map and app update instructions')
p('Project root on this workstation')
p(str(ROOT)).runs[0].font.size=Pt(9)
p('Use relative paths for repository files. The app prepends its deployment base, so the same records can work locally and later on GitHub Pages. Keep original source files and stable equipment/setup IDs.')
table(['Content','Directory or file','How it is updated'],[
('Equipment records','public/library/equipment.json','Admin on localhost; export and commit on static hosting'),('Equipment pictures','public/library/images/','Admin upload or copy file and record image path'),('Viewer models','public/library/models/','Admin GLB upload; record model3d path'),('CAD and model references','public/library/references/<equipment-id>/','Copy source files; record references and permissions'),('Catalog setup photos','public/assets/','Copy photos; register filenames and source markers'),('Shared manuals and SOPs','public/documents/','Copy files; register document metadata in catalog'),('Setup and guide definitions','src/data/catalog.js','Update setup metadata and approved guide content'),('Saved setup layouts','public/library/setups/  proposed','Export Build setup JSON and place here; not yet an automatic catalog importer'),('Private calibration / serial data','Owner approved restricted drive','Record controlled location; no public copy by default')],[1.6,2.65,2.85],9.5)
h2('Record the actual file locations')
table(['Content ID','Original file or controlled location','Target relative path','Owner approval'],[(blank,blank,blank,blank) for _ in range(2)],[1.0,2.3,2.65,1.15],10)
p('Browser-attached documents and drafts stay local to the browser origin. Export layouts and provide referenced files. Serials, calibration, per-port limits and approvals may require app schema extensions; these fields are not all supported in Admin yet.')

page('Review and handover checklist')
checklist(['All equipment names and variants match the installed units','Specifications include units evidence dates and local validated limits','Equipment photos and models are correctly associated with stable IDs','Every physical path and both connector ends are recorded','Setup photos marker locations and exported layout JSON agree','Optical electrical control and observation roles are distinguished','Guidelines have a responsible reviewer and clear version','Measurement limits and reference methods have been validated','Timing assumptions are supported by real measurement records','Files are readable and all referenced directory paths resolve','Publication permissions and controlled records are separated','Remaining unknowns are listed rather than replaced with guesses'])
table(['Open item','Owner','Evidence needed','State / due date'],[(blank,blank,blank,blank) for _ in range(4)],[2.55,1.1,2.2,1.25],10)
field('Completed by date');field('Technical reviewer and approval scope');field('Files provided with this workbook')
p('Handover: provide the completed DOCX, supporting manuals/photos/CAD and exported setup JSON. The next app update should map these records to the directories above, retain source references, identify unsupported fields that need implementation, and report unresolved inconsistencies for review. Completing this workbook does not automatically synchronise the app.')

page('Starting equipment evidence and known gaps')
p('The following values are copied from the current app as a starting reference. They are not newly validated by this workbook. Replace or annotate them with the installed configuration and your evidence.')
for i,e in enumerate(catalog['equipment']):
    if i and i%3==0:page('Starting equipment evidence continued')
    h2(e['id']+'  '+e['name'])
    p(e['model']+'  |  '+e['location']+'  |  '+e['status'])
    p('; '.join(f'{a}: {b}' for a,b in e['specs']))
    p(e['note'])
    if e.get('url'):p('Source: '+e['url']).runs[0].font.size=Pt(9)
    field('Correction evidence and reviewer',lines=2)

out=ROOT/'public/documents/CORNERSTONE_Setup_Collection_Workbook.docx';out.parent.mkdir(exist_ok=True,parents=True)
doc.core_properties.title='CORNERSTONE testing setup collection workbook'
doc.core_properties.subject='Equipment models actual measurement setups work guidelines and directory mapping'
doc.core_properties.author='CORNERSTONE'
for element in list(doc.styles.element.xpath('.//w:pBdr')) + list(doc.element.xpath('.//w:pBdr')):
    element.getparent().remove(element)
doc.save(out);print(out)

