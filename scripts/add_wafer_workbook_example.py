from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
target = ROOT/'public/documents/CORNERSTONE_Setup_Collection_Workbook.docx'
backup = ROOT/'.local/CORNERSTONE_Setup_Collection_Workbook_v1_0.docx'
if not backup.exists(): copy2(target, backup)
doc = Document(backup)
helpers = (ROOT/'scripts/create_collection_workbook.py').read_text(encoding='utf-8')
exec(helpers[helpers.index('def p('):helpers.index("blank='[Enter / verify]'")])

def value(label, text):
    para=p();para.add_run(label+'  ').bold=True;para.add_run(text)

def link(label,url):
    para=p();node=OxmlElement('w:hyperlink')
    node.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'243650');props.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
    run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);node.append(run);para._p.append(node)

# Keep the blank collection forms and current app evidence intact.
intro = doc.paragraphs[2]
intro.add_run(' A filled wafer-scale example is included at the end. Start there to see how to complete the blank sheets.').bold=True
for para in doc.sections[0].footer.paragraphs:
    for run in para.runs:
        run.text=run.text.replace('Version 1.0','Version 1.1').replace('30 September 2026','1 October 2026')

page('Filled wafer scale equipment register')
p('Use this worked example to fill the earlier blank sheets. Equipment names and the new / old grouping come from your wafer-scale tester list dated 1 October 2026. Manufacturer information has its own source reference. Physical installation, serials, fitted options, calibration and bench connections still require your inspection.')
p('The IDs below are proposed stable records. Keep separate entries for each physical unit. The legacy models already appear in the app for chip testing; that does not confirm that the same physical units belong to the wafer tester.')
table(['Proposed equipment ID','Model and role','Starting specification and evidence'],[
('wst-laser-new','Keysight N7776C\nNew tunable laser','Range depends on option; bidirectional sweep up to 200 nm/s. Typical wavelength accuracy ±1.5 pm. [S1]'),
('wst-head-interface','Keysight N7749C\nNew head interface','Controls two or four 8162-C heads depending on configuration; LAN / USB readout and SCPI. [S2]'),
('wst-power-head','Keysight 8162-C family\nNew optical power head','Exact model suffix unknown. Head and adapter determine wavelength, power limits and connector compatibility. [S3]'),
('wst-laser-old','Keysight 81940A\nOld tunable laser','1520–1630 nm; +13 dBm output; wavelength accuracy ±20 pm. Requires compatible 816x mainframe. [S4]'),
('wst-sensor-old','Keysight 81634B\nOld optical sensor','Published power uncertainty ±2.5%; sensitivity −110 dBm; low polarisation dependence <±0.005 dB. [S5]'),
('wst-picoammeter','Keithley 6487\nPicoammeter / voltage source','Reference manual 6487-901-01 revision D. Enter measurement ranges and compliance from the applicable manual and local recipe. [S6]'),
('wst-power-supply','Keysight E3640A\nElectrical power supply','Single output, 30 W; 0–8 V / 3 A OR 0–20 V / 1.5 A. These are alternative ranges. [S7]')
],[1.45,2.05,3.6],10)
p('State for all seven records: Collected / needs local verification. Manufacturer ratings are not the approved DUT operating limits.')

page('Filled equipment identity and specification example')
value('Equipment ID','wst-laser-new')
value('Name and purpose','New wafer-scale tunable laser source for optical wavelength sweeps')
value('Manufacturer and model','Keysight N7776C; fitted wavelength / power option [Enter from label or instrument configuration]')
value('Serial asset tag quantity','[Enter exact unit details; do not use the model number as the serial]')
value('Location owner calibration','Wafer-scale tester; exact room, owner and calibration due date [Enter / verify]')
table(['Parameter','Manufacturer reference','Local validated value'],[
('Wavelength range','Options cover 1240–1380, 1340–1495, 1450–1650 or 1490–1640 nm. [S1]','[Enter fitted option and usable range]'),
('Sweep speed','Up to 200 nm/s, bidirectional. [S1]','[Enter tested speed and trigger conditions]'),
('Wavelength accuracy','Typical ±1.5 pm, including continuous sweep. [S1]','[Enter verification result and date]'),
('Output power','Option dependent; use the fitted option specification. [S1]','[Enter setpoint and approved DUT limit]'),
('Output connector','[Check installed connector and adapter]','[Record exact port and connector]')
],[1.5,3.35,2.25],10)
value('Evidence and review date','Manufacturer page S1 checked 1 October 2026; installed configuration review [Enter name and date]')
value('Firmware software and control','[Enter firmware, driver, acquisition software and trigger arrangement]')
p('Example of a completed local field: write the measured or observed value, its unit, the evidence filename and the review date. Until the unit has been checked, keep the bracketed prompt or write Unknown. Do not select a wavelength option simply because the setup is called C-band.')

page('Filled model and picture collection example')
value('Equipment ID','wst-laser-new')
value('Model target','Detailed representation of the installed Keysight N7776C variant')
value('Dimensions and moving parts','[Measure width depth height in mm; record origin, axes, removable parts and cable clearance]')
table(['Asset','Current evidence','Proposed app path'],[
('Front picture','Not supplied; take a photo of the actual unit','library/images/wst-laser-new-front.jpg'),
('Viewer model','Not supplied; obtain CAD or build from photos and dimensions','library/models/wst-laser-new.glb'),
('Original CAD and drawings','[Enter original filename, revision and source]','library/references/wst-laser-new/'),
('Port closeups and other views','[Enter front rear side and top filenames]','library/references/wst-laser-new/')
],[1.4,3.0,2.7],10)
p('These are target locations, not claims that the files already exist. Physical project folders begin with public/, for example public/library/models/wst-laser-new.glb. The app record uses library/models/wst-laser-new.glb.')
value('Licence and publication permission','[Enter vendor / owner permission before sharing CAD or photos in a public repository]')
value('Viewer checks','Self-contained GLB 2.0 with embedded textures, at most 40 MB; check orientation, appearance and the exact equipment variant')
value('Acceptance and review','[Enter reviewer, date, dimensions used and remaining inaccuracies]')
p('Repeat this sheet for the N7749C, the exact 8162-C head, the 81940A assembly and mainframe, the 81634B assembly, the 6487 and the E3640A. The current viewer normalises instrument size for illustration; physical scale, animated mechanisms and port snapping need separate implementation.')

page('Filled actual wafer scale setup example')
value('Setup name and proposed IDs','Wafer-scale tester; wst-optical-new, wst-optical-old and wst-electrical')
value('Scale and modes','Wafer; optical and electrical. Exact wafer platform and operating readiness [Enter / verify]')
value('Optical band and DUT','[Enter fitted laser range, head range, fibre range, wafer size, coupling and DUT holder]')
value('Overview photo','[Insert annotated bench photo; give filename, date and marker positions]')
table(['Configuration','Equipment supplied in your list','Still needed'],[
('New optical','N7776C; N7749C; 8162-C head family','Laser option; exact head model and count; ports, adapters, fibres, polarisation and DUT routing'),
('Old optical','81940A; 81634B','Mainframe model; unit identity; ports, connectors, fibres and actual legacy wiring'),
('Electrical','Keithley 6487; Keysight E3640A','Each instrument role; probes and DUT pads; return / grounding; compliance and sequencing')
],[1.15,2.55,3.4],10)
h2('Connection sheet example to complete at the bench')
p('Suggested rows below are questions for tracing the setup, not confirmed wiring. Give each link its own ID. Record the connector at each cable end separately.')
table(['Link ID and type','From and to','Details to enter'],[
('O1 optical','N7776C output → [next actual component / port]','Fibre manufacturer and part number; SM / PM; range; connector A and B; length'),
('O2 optical','[DUT output / intervening component] → exact 8162-C head input','Identify head variant and mating adapter; record fibre and both cable ends'),
('H1 head interface','8162-C head → N7749C head port','Head cable model and port number; this is the head interface link, not an optical fibre'),
('E1 electrical','6487 → [probe / pad / return]','Instrument function, cable, polarity, voltage source setting and current compliance'),
('E2 electrical','E3640A → [actual load / pads]','Load role; range; setpoint; current limit; return / ground'),
('C1 control','Control computer → [each instrument]','LAN USB GPIB or other actual interface; software and trigger connections')
],[1.25,2.65,3.2],10)

page('Filled guideline and directory handover example')
value('Guideline title','Wafer-scale optical and electrical measurement')
value('Owner version and approver','[Enter responsible person, SOP revision, approved date and applicable setup IDs]')
p('Use the stages below to collect the actual lab procedure. Replace each prompt with approved actions and expected results; a vendor product tour is useful training material but is not the local operating procedure.')
table(['Stage','What a completed entry should contain'],[
('Preparation','Authorisation, DUT ID and wafer map, approved setup configuration, hazards and stop conditions'),
('Instrument checks','Installed models and options, calibration, software and required warm-up / checks from the SOP'),
('Connect and align','Confirmed connection IDs, fibre model and both connectors, reference and alignment acceptance criteria'),
('Set and acquire','Actual wavelength / electrical sequence, power, averaging, compliance, repeats and filenames'),
('Review and shut down','Saturation / drift checks, approved recovery, unload and shutdown steps'),
('Post process','Processing script / version, units, reference subtraction, plots, acceptance and report location')
],[1.45,5.65],10)
value('Timing evidence','[Enter preparation, per-wafer, per-device, alignment and per-sweep durations with counts and measured job total]')
h2('Directory handover example')
value('Equipment metadata','public/library/equipment.json after the records are reviewed and imported through Admin')
value('Photos models and references','Use the paths in the filled model sheet, with one stable folder / filename per equipment ID')
value('Guideline target','Proposed public/documents/wst-measurement-guide.pdf or .docx; add document metadata after approval')
value('Layout target','Proposed public/library/setups/wst-optical-new.json from Build setup export; catalog import remains a separate update')
value('Original files','[Enter actual drive path or vendor URL, source revision and owner permission for every attachment]')
p('Next handover: send the completed sheets, referenced files and exported layout JSON. The directory map can then be updated against the actual filenames. Filling this Word document does not automatically change the equipment library or setup catalog.')

page('Wafer scale source references')
p('Manufacturer pages checked 1 October 2026. Your list establishes the association with the wafer-scale tester; these sources establish product information. Installed options, physical connections and operating readiness remain bench checks.')
sources=[
('S1 N7776C tunable laser source','https://www.keysight.com/us/en/product/N7776C/tunable-laser-source-high-power-lowest-sse-top-line.html'),
('S2 N7749C optical head interface','https://www.keysight.com/us/en/product/N7749C/optical-head-interface.html'),
('S3 8162-C and N7749C head family data sheet','https://www.keysight.com/us/en/assets/3121-1226/data-sheets/8162-C-and-N7749C-Optical-Power-Meter-Heads-and-Optical-Head-Interface.pdf'),
('S4 81940A legacy tunable laser','https://www.keysight.com/us/en/product/81940A/compact-tunable-laser-source-continuous-sweep-mode-1520nm-1630nm.html#resources'),
('S5 81634B legacy optical power sensor','https://www.keysight.com/us/en/product/81634B/low-polarization-dependence-optical-power-sensor.html'),
('S6 Keithley 6487 reference manual','https://www.tek.com/en/specialty-instruments/keithley-series-6400-picoammeters-manual/model-6487-picoammeter-voltage-source-reference-manual'),
('S7 E3640A power supply','https://www.keysight.com/us/en/product/E3640A/30w-power-supply-8v-3a-20v-1-5a.html'),
('Further reading Optical power measurement products','https://www.keysight.com/us/en/products/photonic-products/optical-component-test-products/optical-power-measurement-products.html'),
('Training video Keysight swept wavelength measurement tour','https://www.keysight.com/us/en/assets/3122-1541/product-tours/Lambda-Scan-Solutions-Swept-Wavelength-with-Polarization-Dependence-and-Alignment.mp4')]
for label,url in sources:link(label,url)
p('The training video link is retained from your list as supplementary material. It has not been used as evidence of the actual wafer-tester wiring or as an approved lab guideline.')
p('For S6, the reference page lists manual 6487-901-01 revision D released 28 October 2020. Confirm that this revision applies to the installed unit. Retain links or controlled local copies according to the source permissions.')
doc.core_properties.modified = __import__('datetime').datetime(2026,10,1,12,0,0)
doc.core_properties.version='1.1'
doc.save(target)
print(target)
